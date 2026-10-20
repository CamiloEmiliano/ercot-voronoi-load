from datetime import date, datetime

import numpy as np
import openpyxl
import pandas as pd

from scripts.normalize_ercot_fuel_mix_2018 import (
    DEFAULT_INPUT,
    _interval_minutes,
    _sort_records,
    normalize_2018,
    normalize_row,
)


def _headers():
    return [15 * (index + 1) for index in range(96)]


def test_spring_forward_preserves_2018_0215_through_0300_gap():
    values = [float(index) for index in range(96)]
    values[8:12] = [None] * 4
    records = normalize_row(date(2018, 3, 11), "Biomass", values, _headers())

    assert len(records) == 96
    assert [records[index]["timestamp_local"] for index in range(7, 13)] == [
        "2018-03-11T02:00:00-06:00",
        "2018-03-11T02:15:00-06:00",
        "2018-03-11T02:30:00-06:00",
        "2018-03-11T02:45:00-06:00",
        "2018-03-11T03:00:00-06:00",
        "2018-03-11T03:15:00-05:00",
    ]
    assert all(np.isnan(records[index]["generation_mwh"]) for index in range(8, 12))
    assert records[12]["generation_mwh"] == 12.0


def test_fall_back_preserves_both_0115_through_0200_passes():
    values = [float(index) for index in range(96)]
    dst_values = [1000.0 + index for index in range(4)]
    records = normalize_row(date(2018, 11, 4), "Biomass", values, _headers(), dst_values)

    assert len(records) == 100
    assert [records[index]["timestamp_local"] for index in range(4, 8)] == [
        "2018-11-04T01:15:00-05:00",
        "2018-11-04T01:30:00-05:00",
        "2018-11-04T01:45:00-05:00",
        "2018-11-04T02:00:00-05:00",
    ]
    assert [records[index]["timestamp_local"] for index in range(96, 100)] == [
        "2018-11-04T01:15:00-06:00",
        "2018-11-04T01:30:00-06:00",
        "2018-11-04T01:45:00-06:00",
        "2018-11-04T02:00:00-06:00",
    ]
    assert [record["generation_mwh"] for record in records[96:100]] == dst_values


def test_spring_placeholders_sort_as_timestamp_major_contiguous_block():
    values = [float(index) for index in range(96)]
    values[8:12] = [None] * 4
    records = [
        record
        for fuel in ("Biomass", "Coal")
        for record in normalize_row(date(2018, 3, 11), fuel, values, _headers())
    ]
    frame = _sort_records(pd.DataFrame.from_records(records))
    null_positions = np.flatnonzero(frame["generation_mwh"].isna().to_numpy())

    assert len(null_positions) == 8
    assert np.array_equal(null_positions, np.arange(null_positions[0], null_positions[-1] + 1))
    gap = frame.iloc[null_positions]
    assert gap["fuel_type"].tolist() == ["Biomass", "Coal"] * 4
    assert gap["timestamp_local"].tolist() == [
        f"2018-03-11T{hour:02d}:{minute:02d}:00-06:00"
        for hour, minute in ((2, 15), (2, 15), (2, 30), (2, 30), (2, 45), (2, 45), (3, 0), (3, 0))
    ]
    assert frame.iloc[null_positions[0] - 1]["timestamp_local"] == "2018-03-11T02:00:00-06:00"
    assert frame.iloc[null_positions[-1] + 1]["timestamp_local"] == "2018-03-11T03:15:00-05:00"


def test_text_and_numeric_headers_are_supported():
    assert _interval_minutes("0:15") == 15
    assert _interval_minutes("0:00") == 1440
    assert _interval_minutes(15 / 1440) == 15


def test_generated_parquet_persists_and_matches_dst_source_cells(tmp_path):
    output = tmp_path / "fuel_mix_2018.parquet"
    frame = normalize_2018(output_path=output, workers=1)
    saved = pd.read_parquet(output)

    assert frame.attrs["sampled_values_verified"] == 120
    assert saved.columns.tolist() == ["timestamp_local", "fuel_type", "generation_mwh"]
    assert len(saved) == 315396
    assert saved.groupby("fuel_type").size().eq(35044).all()

    march = saved[saved["timestamp_local"].str.startswith("2018-03-11T")].reset_index(drop=True)
    null_positions = np.flatnonzero(march["generation_mwh"].isna().to_numpy())
    assert len(null_positions) == 36
    assert np.array_equal(null_positions, np.arange(null_positions[0], null_positions[-1] + 1))
    assert march.iloc[null_positions[0] - 1]["timestamp_local"] == "2018-03-11T02:00:00-06:00"
    assert march.iloc[null_positions[-1] + 1]["timestamp_local"] == "2018-03-11T03:15:00-05:00"

    workbook = openpyxl.load_workbook(DEFAULT_INPUT, read_only=True, data_only=True)
    try:
        sheet = workbook["Nov"]
        sheet.reset_dimensions()
        rows = [
            row for row in sheet.iter_rows(values_only=True)
            if isinstance(row[0], datetime) and row[0].date() == date(2018, 11, 4)
        ]
        assert len(rows) == 9
        for row in rows:
            fuel_type = row[1]
            daily = saved[
                (saved["fuel_type"] == fuel_type)
                & saved["timestamp_local"].str.startswith("2018-11-04T")
            ]
            for source_index, local_time in enumerate(("01:15", "01:30", "01:45", "02:00")):
                first = daily[daily["timestamp_local"] == f"2018-11-04T{local_time}:00-05:00"]
                second = daily[daily["timestamp_local"] == f"2018-11-04T{local_time}:00-06:00"]
                assert first["generation_mwh"].item() == float(row[8 + source_index])
                assert second["generation_mwh"].item() == float(row[100 + source_index])
    finally:
        workbook.close()