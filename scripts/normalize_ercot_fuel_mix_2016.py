"""Normalize the 2016 ERCOT Fuel Mix workbook into interval-level Parquet."""

from __future__ import annotations

import argparse
from collections.abc import Sequence
from concurrent.futures import ProcessPoolExecutor
from datetime import date, datetime, time, timedelta, timezone
import json
import math
from pathlib import Path
import re
from typing import Any
from zoneinfo import ZoneInfo

import pandas as pd
import openpyxl

if __package__:
    from .ercot_fuel_mix_qc import write_parquet_with_qc
else:
    from ercot_fuel_mix_qc import write_parquet_with_qc

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = PROJECT_ROOT / "data/raw/ercot_fuel_mix/IntGenbyFuel2016.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data/interim/fuel_mix/fuel_mix_2016.parquet"
WORKSHEETS = (
    "Jan16", "Feb16", "Mar16", "Apr16", "May16", "Jun16",
    "Jul16", "Aug16", "Sep16", "Oct16", "Nov16", "Dec16",
)
FUEL_ROW_PATTERN = re.compile(r"^(?P<date>\d{2}/\d{2}/\d{2})_(?P<fuel>.+)$")
CENTRAL = ZoneInfo("America/Chicago")
OUTPUT_COLUMNS = ("timestamp_local", "fuel_type", "generation_mwh")


def _parse_date_fuel(value: Any) -> tuple[date, str] | None:
    if not isinstance(value, str) or not value.strip():
        return None
    match = FUEL_ROW_PATTERN.fullmatch(value.strip())
    if match is None:
        return None
    return datetime.strptime(match.group("date"), "%m/%d/%y").date(), match.group("fuel")


def _interval_minutes(header: Any) -> int:
    if isinstance(header, timedelta) and header == timedelta(0):
        return 1440
    if isinstance(header, datetime) and header.date() == date(1900, 1, 1) and header.time() == time():
        return 1440
    if isinstance(header, time):
        minutes = header.hour * 60 + header.minute
        if minutes == 0:
            return 1440
        if header.second or header.microsecond:
            raise ValueError(f"Invalid interval header: {header!r}")
        return minutes
    if isinstance(header, (int, float)):
        minutes = int(round(float(header) * 24 * 60))
        if minutes == 0:
            return 1440
        if 0 < minutes <= 1440 and minutes % 15 == 0:
            return minutes
    if isinstance(header, str):
        parts = header.split(":")
        if len(parts) == 2:
            hour, minute = (int(part) for part in parts)
            minutes = hour * 60 + minute
            if 0 < minutes <= 1440 and minutes % 15 == 0:
                return minutes
    raise ValueError(f"Invalid 15-minute interval header: {header!r}")


def _energy_value(value: Any) -> float:
    if value is None or value == "":
        return math.nan
    return float(value)


def _serialize_local(local_end: datetime, fold: int | None = 0) -> str:
    if fold is None:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
    return local_end.replace(tzinfo=CENTRAL, fold=fold).isoformat(timespec="seconds")


def normalize_row(
    operating_date: date,
    fuel_type: str,
    interval_values: Sequence[Any],
    interval_headers: Sequence[int],
    dst_values: Sequence[Any] = (),
) -> list[dict[str, Any]]:
    """Map one 2016 daily fuel row to its 96 slots and optional second pass."""
    spring_forward = operating_date == date(2016, 3, 13)
    fall_back = operating_date == date(2016, 11, 6)
    if len(interval_values) != 96 or len(interval_headers) != 96:
        raise ValueError("2016 rows require exactly 96 regular interval values and headers")

    records: list[dict[str, Any]] = []
    for slot, endpoint_minutes in enumerate(interval_headers):
        local_end = datetime.combine(operating_date, time()) + timedelta(minutes=endpoint_minutes)
        is_spring_gap = spring_forward and slot in range(7, 11)
        timestamp_local = _serialize_local(local_end, fold=None if is_spring_gap else 0)
        energy = math.nan if is_spring_gap else _energy_value(interval_values[slot])
        records.append({
            "timestamp_local": timestamp_local,
            "fuel_type": fuel_type,
            "generation_mwh": energy,
        })

    if fall_back:
        if len(dst_values) != 4:
            raise ValueError("2016-11-06 rows require four second-pass DST values")
        for dst_index in range(4):
            local_end = datetime.combine(operating_date, time()) + timedelta(minutes=60 + 15 * dst_index)
            records.append({
                "timestamp_local": _serialize_local(local_end, fold=1),
                "fuel_type": fuel_type,
                "generation_mwh": _energy_value(dst_values[dst_index]),
            })
    return records


def _normalize_sheet(args: tuple[str, str]) -> list[dict[str, Any]]:
    workbook_path, sheet_name = args
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        interval_headers = [_interval_minutes(value) for value in headers[2:98]]
        if interval_headers != list(range(15, 1441, 15)):
            raise ValueError(f"{sheet_name}: interval headers do not match quarter-hour endpoints")

        records: list[dict[str, Any]] = []
        for row_number, row in enumerate(rows, start=2):
            parsed = _parse_date_fuel(row[0] if row else None)
            if parsed is None:
                continue
            operating_date, fuel_type = parsed
            if len(row) < 98:
                raise ValueError(f"{sheet_name} row {row_number}: fewer than 96 interval cells")
            interval_values = row[2:98]
            if operating_date == date(2016, 3, 13) and any(value is not None for value in interval_values[7:11]):
                raise ValueError(f"{sheet_name} row {row_number}: expected J:M spring structural blanks")
            dst_values = row[98:102] if operating_date == date(2016, 11, 6) else ()
            records.extend(normalize_row(operating_date, fuel_type, interval_values, interval_headers, dst_values))
        return records
    finally:
        workbook.close()


def _sort_records(frame: pd.DataFrame) -> pd.DataFrame:
    sort_utc = pd.to_datetime(frame["timestamp_local"], format="ISO8601", utc=True)
    spring_gap = frame["timestamp_local"].str.startswith("2016-03-13T02:") & frame["generation_mwh"].isna()
    # Structural placeholders have no real instant; keep them together before 03:00.
    sort_utc.loc[spring_gap] = pd.Timestamp("2016-03-13T07:59:59Z")
    return (
        frame.assign(_sort_utc=sort_utc)
        .sort_values(["_sort_utc", "timestamp_local", "fuel_type"], ignore_index=True)
        .drop(columns="_sort_utc")
    )


def normalize_2016(
    input_path: Path = DEFAULT_INPUT,
    output_path: Path = DEFAULT_OUTPUT,
    workers: int | None = None,
) -> pd.DataFrame:
    """Create one long-form 2016 Parquet artifact from the twelve monthly sheets."""
    if not input_path.is_file():
        raise FileNotFoundError(input_path)
    jobs = [(str(input_path), sheet_name) for sheet_name in WORKSHEETS]
    with ProcessPoolExecutor(max_workers=workers) as executor:
        partitions = list(executor.map(_normalize_sheet, jobs))
    records = [record for partition in partitions for record in partition]
    frame = pd.DataFrame.from_records(records, columns=OUTPUT_COLUMNS)
    frame["generation_mwh"] = frame["generation_mwh"].astype("float64")
    frame = _sort_records(frame)
    frame.attrs["sampled_values_verified"] = write_parquet_with_qc(frame, 2016, input_path, output_path)
    return frame


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--workers", type=int, default=None)
    args = parser.parse_args()
    frame = normalize_2016(args.input, args.output, args.workers)
    print(json.dumps({
        "output": str(args.output),
        "rows": len(frame),
        "fuels": sorted(frame["fuel_type"].unique().tolist()),
        "generation_nan": int(frame["generation_mwh"].isna().sum()),
        "sampled_values_verified": frame.attrs["sampled_values_verified"],
    }, indent=2))


if __name__ == "__main__":
    main()