"""Independent workbook-to-Parquet sampled quality checks for annual ERCOT files."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
import math
from pathlib import Path
import random
import re
import os
import tempfile
from typing import Any, Iterator
from zoneinfo import ZoneInfo

import openpyxl
import pandas as pd
import xlrd
from openpyxl.utils import get_column_letter

CENTRAL = ZoneInfo("America/Chicago")
MONTH_SHEETS = {
    2014: tuple(f"{month}14" for month in ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")),
    2015: tuple(f"{month}15" for month in ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")),
    2016: tuple(f"{month}16" for month in ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")),
    2017: ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
    2018: ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
    2019: ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
    2020: ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
    2021: ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
}


@dataclass(frozen=True)
class SourceValue:
    month: int
    cell: str
    timestamp_local: str
    fuel_type: str
    expected: float


def _interval_minutes(header: Any) -> int:
    if isinstance(header, datetime):
        minutes = header.hour * 60 + header.minute
    elif isinstance(header, time):
        minutes = header.hour * 60 + header.minute
    elif isinstance(header, timedelta):
        minutes = int(header.total_seconds() // 60)
    elif isinstance(header, (int, float)):
        minutes = int(round(float(header) * 24 * 60))
    elif isinstance(header, str):
        parts = header.split(":")
        if len(parts) != 2:
            raise ValueError(f"Invalid source interval header: {header!r}")
        hour, minute = (int(part) for part in parts)
        minutes = hour * 60 + minute
    else:
        raise ValueError(f"Invalid source interval header: {header!r}")
    return 1440 if minutes == 0 else minutes


def _parse_date_fuel(value: Any) -> tuple[date, str] | None:
    if not isinstance(value, str):
        return None
    match = re.fullmatch(r"\s*(\d{2}/\d{2}/\d{2})_(.+?)\s*", value)
    if match is None:
        return None
    return datetime.strptime(match.group(1), "%m/%d/%y").date(), match.group(2)


def _source_date_fuel(year: int, row: tuple[Any, ...]) -> tuple[date, str] | None:
    if year in (2017, 2018, 2019, 2020, 2021):
        raw_date = row[0] if row else None
        fuel = row[1] if len(row) > 1 else None
        if isinstance(raw_date, datetime):
            raw_date = raw_date.date()
        if isinstance(raw_date, date) and isinstance(fuel, str):
            return raw_date, fuel
        return None
    return _parse_date_fuel(row[0] if row else None)


def _timestamp(year: int, operating_date: date, endpoint: int, slot: int, *, second_pass: bool = False) -> str:
    local_end = datetime.combine(operating_date, time()) + timedelta(minutes=endpoint)
    if second_pass:
        offset = timedelta(hours=-6)
        return local_end.replace(tzinfo=timezone(offset)).isoformat(timespec="seconds")
    if year == 2014 and operating_date == date(2014, 11, 2) and 3 <= slot <= 6:
        return local_end.replace(tzinfo=CENTRAL, fold=1).isoformat(timespec="seconds")
    if year == 2017 and operating_date == date(2017, 11, 5) and slot == 7:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-5))).isoformat(timespec="seconds")
    if year == 2018 and operating_date == date(2018, 11, 4) and 4 <= slot <= 7:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-5))).isoformat(timespec="seconds")
    if year == 2019 and operating_date == date(2019, 11, 3) and 4 <= slot <= 7:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-5))).isoformat(timespec="seconds")
    if year == 2020 and operating_date == date(2020, 11, 1) and 4 <= slot <= 7:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-5))).isoformat(timespec="seconds")
    if year == 2021 and operating_date == date(2021, 11, 7) and 4 <= slot <= 7:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-5))).isoformat(timespec="seconds")
    if year in (2014, 2015) and operating_date in (date(year, 3, 9) if year == 2014 else date(year, 3, 8),):
        if year == 2014 and 7 <= slot <= 10:
            return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
        if year == 2015 and 7 <= slot <= 10:
            return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
    if year == 2016 and operating_date == date(2016, 3, 13) and 7 <= slot <= 10:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
    if year == 2017 and operating_date == date(2017, 3, 12) and 8 <= slot <= 11:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
    if year == 2018 and operating_date == date(2018, 3, 11) and 8 <= slot <= 11:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
    if year == 2019 and operating_date == date(2019, 3, 10) and 8 <= slot <= 11:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
    if year == 2020 and operating_date == date(2020, 3, 8) and 8 <= slot <= 11:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
    if year == 2021 and operating_date == date(2021, 3, 14) and 8 <= slot <= 11:
        return local_end.replace(tzinfo=timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
    return local_end.replace(tzinfo=CENTRAL, fold=0).isoformat(timespec="seconds")


def _is_value(value: Any) -> bool:
    return value is not None and value != ""


def _expected_value(year: int, value: Any) -> float:
    if value == "-" and year in (2014, 2015):
        return 0.0
    return float(value)


def _candidate_for_base(
    year: int,
    month: int,
    row_number: int,
    row: tuple[Any, ...],
    operating_date: date,
    fuel_type: str,
    source_index: int,
    source_value: Any,
    headers: list[int],
    source_column_start: int,
) -> SourceValue | None:
    if not _is_value(source_value):
        return None

    target_slot = source_index
    if year == 2014 and operating_date == date(2014, 3, 9):
        if source_index >= 92:
            return None
        target_slot = source_index if source_index < 7 else source_index + 4
    elif year == 2015 and operating_date == date(2015, 3, 8):
        if source_index >= 92:
            return None
        target_slot = source_index if source_index < 7 else source_index + 4

    timestamp_local = _timestamp(year, operating_date, headers[target_slot], target_slot)
    return SourceValue(
        month=month,
        cell=f"{get_column_letter(source_column_start + source_index)}{row_number}",
        timestamp_local=timestamp_local,
        fuel_type=fuel_type,
        expected=_expected_value(year, source_value),
    )


def _iter_xls_candidates(year: int, workbook_path: Path) -> Iterator[SourceValue]:
    workbook = xlrd.open_workbook(str(workbook_path), on_demand=True)
    try:
        for month, sheet_name in enumerate(MONTH_SHEETS[year], start=1):
            sheet = workbook.sheet_by_name(sheet_name)
            headers = [_interval_minutes(value) for value in sheet.row_values(0)[2:98]]
            for row_index in range(1, sheet.nrows):
                row = tuple(sheet.row_values(row_index))
                parsed = _parse_date_fuel(row[0] if row else None)
                if parsed is None:
                    continue
                operating_date, fuel_type = parsed
                for source_index, value in enumerate(row[2:98]):
                    target_slot = source_index
                    if year == 2014 and operating_date == date(2014, 11, 2) and 3 <= source_index <= 6:
                        target_slot = source_index
                        stamp = _timestamp(year, operating_date, headers[target_slot], target_slot, second_pass=True)
                        candidate = SourceValue(month, f"{get_column_letter(3 + source_index)}{row_index + 1}", stamp, fuel_type, _expected_value(year, value)) if _is_value(value) else None
                    else:
                        candidate = _candidate_for_base(year, month, row_index + 1, row, operating_date, fuel_type, source_index, value, headers, 3)
                    if candidate is not None:
                        yield candidate
                if year == 2015 and operating_date == date(2015, 11, 1):
                    for dst_index, value in enumerate(row[98:102]):
                        if not _is_value(value):
                            continue
                        endpoint = 60 + 15 * dst_index
                        stamp = _timestamp(year, operating_date, endpoint, endpoint // 15 - 1, second_pass=True)
                        yield SourceValue(month, f"{get_column_letter(99 + dst_index)}{row_index + 1}", stamp, fuel_type, _expected_value(year, value))
    finally:
        workbook.release_resources()


def _iter_xlsx_candidates(year: int, workbook_path: Path) -> Iterator[SourceValue]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        for month, sheet_name in enumerate(MONTH_SHEETS[year], start=1):
            sheet = workbook[sheet_name]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows)
            source_column_start = 5 if year in (2017, 2018, 2019, 2020, 2021) else 3
            headers = [_interval_minutes(value) for value in header_row[source_column_start - 1:source_column_start - 1 + 96]]
            for row_number, row in enumerate(rows, start=2):
                parsed = _source_date_fuel(year, row)
                if parsed is None:
                    continue
                operating_date, fuel_type = parsed
                values_start = source_column_start - 1
                interval_values = row[values_start:values_start + 96]
                for source_index, value in enumerate(interval_values):
                    candidate = _candidate_for_base(
                        year, month, row_number, row, operating_date, fuel_type,
                        source_index, value, headers, source_column_start,
                    )
                    if candidate is not None:
                        yield candidate
                if year in (2016, 2017, 2018, 2019, 2020, 2021) and operating_date == date(year, 11, 6 if year == 2016 else 5 if year == 2017 else 4 if year == 2018 else 3 if year == 2019 else 1 if year == 2020 else 7):
                    dst_start = 98 if year == 2016 else 100
                    for dst_index, value in enumerate(row[dst_start:dst_start + 4]):
                        if not _is_value(value):
                            continue
                        endpoint = (60 if year == 2016 else 75) + 15 * dst_index
                        target_slot = endpoint // 15 - 1
                        stamp = _timestamp(year, operating_date, endpoint, target_slot, second_pass=True)
                        yield SourceValue(
                            month, f"{get_column_letter(dst_start + 1 + dst_index)}{row_number}",
                            stamp, fuel_type, _expected_value(year, value),
                        )
    finally:
        workbook.close()


def _iter_candidates(year: int, workbook_path: Path) -> Iterator[SourceValue]:
    if year not in MONTH_SHEETS:
        raise ValueError(f"Unsupported ERCOT workbook year: {year}")
    if year in (2014, 2015):
        yield from _iter_xls_candidates(year, workbook_path)
    else:
        yield from _iter_xlsx_candidates(year, workbook_path)


def audit_workbook_to_parquet(
    year: int,
    workbook_path: str | Path,
    parquet_path: str | Path,
    *,
    sample_per_month: int = 10,
    seed: int | None = None,
) -> int:
    """Require exact parity for a reproducible random sample from every month."""
    workbook_path = Path(workbook_path)
    parquet_path = Path(parquet_path)
    if sample_per_month <= 0:
        raise ValueError("sample_per_month must be positive")

    persisted = pd.read_parquet(parquet_path)
    expected_columns = ["timestamp_local", "fuel_type", "generation_mwh"]
    if persisted.columns.tolist() != expected_columns:
        raise AssertionError(f"{parquet_path}: unexpected Parquet columns {persisted.columns.tolist()!r}")
    keys = list(zip(persisted["timestamp_local"], persisted["fuel_type"]))
    if len(set(keys)) != len(keys):
        raise AssertionError(f"{parquet_path}: duplicate (timestamp_local, fuel_type) keys")
    parquet_values = dict(zip(keys, persisted["generation_mwh"]))

    candidates_by_month: dict[int, list[SourceValue]] = {month: [] for month in range(1, 13)}
    for candidate in _iter_candidates(year, workbook_path):
        candidates_by_month[candidate.month].append(candidate)

    rng = random.Random(year if seed is None else seed)
    selected: list[SourceValue] = []
    for month, candidates in candidates_by_month.items():
        if len(candidates) < sample_per_month:
            raise AssertionError(
                f"{year} month {month:02d}: only {len(candidates)} nonblank generation values; "
                f"cannot sample {sample_per_month}"
            )
        selected.extend(rng.sample(candidates, sample_per_month))

    mismatches: list[str] = []
    for sample in selected:
        key = (sample.timestamp_local, sample.fuel_type)
        if key not in parquet_values:
            mismatches.append(
                f"{sample.cell}: missing Parquet key {key!r}; workbook value={sample.expected!r}"
            )
            continue
        actual = parquet_values[key]
        if not math.isfinite(float(actual)) or float(actual) != sample.expected:
            mismatches.append(
                f"{sample.cell} {sample.fuel_type} {sample.timestamp_local}: "
                f"workbook={sample.expected!r}, parquet={actual!r}"
            )
    if mismatches:
        raise AssertionError(
            f"{year} workbook-to-Parquet QC failed: {len(mismatches)}/{len(selected)} sampled values differ:\n"
            + "\n".join(mismatches)
        )
    return len(selected)


def write_parquet_with_qc(
    frame: pd.DataFrame,
    year: int,
    workbook_path: str | Path,
    parquet_path: str | Path,
) -> int:
    """Publish an artifact only after its persisted contents pass the sample audit."""
    parquet_path = Path(parquet_path)
    parquet_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{parquet_path.stem}.",
        suffix=".parquet",
        dir=parquet_path.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)

    try:
        frame.to_parquet(temporary_path, index=False, engine="pyarrow", compression="zstd")
        verified = audit_workbook_to_parquet(year, workbook_path, temporary_path)
        os.replace(temporary_path, parquet_path)
        return verified
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise
