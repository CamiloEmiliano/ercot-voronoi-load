"""Normalize the 2018 ERCOT Fuel Mix workbook into interval-level Parquet."""

from __future__ import annotations

import argparse
from collections.abc import Sequence
from concurrent.futures import ProcessPoolExecutor
from datetime import date, datetime, time, timedelta, timezone
import json
import math
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl
import pandas as pd

if __package__:
    from .ercot_fuel_mix_qc import write_parquet_with_qc
else:
    from ercot_fuel_mix_qc import write_parquet_with_qc

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = PROJECT_ROOT / "data/raw/ercot_fuel_mix/IntGenbyFuel2018.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data/interim/fuel_mix/fuel_mix_2018.parquet"
WORKSHEETS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
FUEL_TYPES = ("Biomass", "Coal", "Gas", "Gas-CC", "Hydro", "Nuclear", "Other", "Solar", "Wind")
CENTRAL = ZoneInfo("America/Chicago")
OUTPUT_COLUMNS = ("timestamp_local", "fuel_type", "generation_mwh")


def _interval_minutes(header: Any) -> int:
    if isinstance(header, datetime):
        minutes = header.hour * 60 + header.minute
    elif isinstance(header, time):
        minutes = header.hour * 60 + header.minute
    elif isinstance(header, timedelta):
        minutes = int(header.total_seconds() // 60)
    elif isinstance(header, (int, float)):
        minutes = int(round(float(header) * 24 * 60))
    elif isinstance(header, str):
        parts = header.split(":")
        if len(parts) != 2:
            raise ValueError(f"Invalid interval header: {header!r}")
        hour, minute = (int(part) for part in parts)
        minutes = hour * 60 + minute
    else:
        raise ValueError(f"Invalid interval header: {header!r}")

    if minutes == 0:
        return 1440
    if not 0 < minutes <= 1440 or minutes % 15:
        raise ValueError(f"Invalid 15-minute interval header: {header!r}")
    return minutes


def _energy_value(value: Any) -> float:
    if value is None or value == "":
        return math.nan
    return float(value)


def _serialize_local(local_end: datetime, utc_offset: timedelta | None = None) -> str:
    timezone_info = timezone(utc_offset) if utc_offset is not None else CENTRAL
    return local_end.replace(tzinfo=timezone_info).isoformat(timespec="seconds")


def normalize_row(
    operating_date: date,
    fuel_type: str,
    interval_values: Sequence[Any],
    interval_headers: Sequence[int],
    dst_values: Sequence[Any] = (),
) -> list[dict[str, Any]]:
    """Map one 2018 daily fuel row, preserving source values and both DST passes."""
    spring_forward = operating_date == date(2018, 3, 11)
    fall_back = operating_date == date(2018, 11, 4)
    if len(interval_values) != 96 or len(interval_headers) != 96:
        raise ValueError("2018 rows require exactly 96 regular interval values and headers")
    if fall_back and len(dst_values) != 4:
        raise ValueError("2018-11-04 rows require four second-pass DST values")

    records: list[dict[str, Any]] = []
    for slot, endpoint_minutes in enumerate(interval_headers):
        local_end = datetime.combine(operating_date, time()) + timedelta(minutes=endpoint_minutes)
        is_spring_gap = spring_forward and slot in range(8, 12)
        first_pass_offset = timedelta(hours=-5) if fall_back and slot in range(4, 8) else None
        records.append({
            "timestamp_local": _serialize_local(
                local_end,
                timedelta(hours=-6) if is_spring_gap else first_pass_offset,
            ),
            "fuel_type": fuel_type,
            "generation_mwh": math.nan if is_spring_gap else _energy_value(interval_values[slot]),
        })

    if fall_back:
        for dst_index, endpoint_minutes in enumerate(range(75, 121, 15)):
            local_end = datetime.combine(operating_date, time()) + timedelta(minutes=endpoint_minutes)
            records.append({
                "timestamp_local": _serialize_local(local_end, timedelta(hours=-6)),
                "fuel_type": fuel_type,
                "generation_mwh": _energy_value(dst_values[dst_index]),
            })
    return records


def _normalize_sheet(args: tuple[str, str]) -> list[dict[str, Any]]:
    workbook_path, sheet_name = args
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        # The source workbook declares A1 as its dimension despite containing data.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        interval_headers = [_interval_minutes(value) for value in headers[4:100]]
        if interval_headers != list(range(15, 1441, 15)):
            raise ValueError(f"{sheet_name}: interval headers do not match quarter-hour endpoints")
        if sheet_name == "Nov" and tuple(headers[100:104]) != (
            "01:15 (DST)", "01:30 (DST)", "01:45 (DST)", "02:00 (DST)"
        ):
            raise ValueError("Nov: expected second-pass DST headers CW:CZ")

        records: list[dict[str, Any]] = []
        for row_number, row in enumerate(rows, start=2):
            operating_date = row[0].date() if row and isinstance(row[0], datetime) else None
            fuel_type = row[1] if len(row) > 1 else None
            if operating_date is None or not isinstance(fuel_type, str) or fuel_type not in FUEL_TYPES:
                continue
            interval_values = row[4:100]
            if len(interval_values) != 96:
                raise ValueError(f"{sheet_name} row {row_number}: fewer than 96 interval cells")
            if operating_date == date(2018, 3, 11) and any(
                value not in (None, "") for value in interval_values[8:12]
            ):
                raise ValueError(f"{sheet_name} row {row_number}: expected M:P spring structural blanks")
            dst_values = row[100:104] if operating_date == date(2018, 11, 4) else ()
            records.extend(normalize_row(operating_date, fuel_type, interval_values, interval_headers, dst_values))
        return records
    finally:
        workbook.close()


def _sort_records(frame: pd.DataFrame) -> pd.DataFrame:
    sort_utc = pd.to_datetime(frame["timestamp_local"], format="ISO8601", utc=True)
    spring_gap = frame["timestamp_local"].str.startswith((
        "2018-03-11T02:15:",
        "2018-03-11T02:30:",
        "2018-03-11T02:45:",
        "2018-03-11T03:00:",
    )) & frame["generation_mwh"].isna()
    # Structural placeholders have no real instant; keep them before 03:15.
    sort_utc.loc[spring_gap] = pd.Timestamp("2018-03-11T08:07:30Z")
    return (
        frame.assign(_sort_utc=sort_utc)
        .sort_values(["_sort_utc", "timestamp_local", "fuel_type"], ignore_index=True)
        .drop(columns="_sort_utc")
    )


def normalize_2018(
    input_path: Path = DEFAULT_INPUT,
    output_path: Path = DEFAULT_OUTPUT,
    workers: int | None = None,
) -> pd.DataFrame:
    """Create one long-form 2018 Parquet artifact and publish only after sampled QC passes."""
    if not input_path.is_file():
        raise FileNotFoundError(input_path)
    jobs = [(str(input_path), sheet_name) for sheet_name in WORKSHEETS]
    with ProcessPoolExecutor(max_workers=workers) as executor:
        partitions = list(executor.map(_normalize_sheet, jobs))
    records = [record for partition in partitions for record in partition]
    frame = pd.DataFrame.from_records(records, columns=OUTPUT_COLUMNS)
    frame["generation_mwh"] = frame["generation_mwh"].astype("float64")
    frame = _sort_records(frame)
    from_qc = write_parquet_with_qc(frame, 2018, input_path, output_path)
    frame.attrs["sampled_values_verified"] = from_qc
    return frame


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--workers", type=int, default=None)
    args = parser.parse_args()
    frame = normalize_2018(args.input, args.output, args.workers)
    print(json.dumps({
        "output": str(args.output),
        "rows": len(frame),
        "fuels": sorted(frame["fuel_type"].unique().tolist()),
        "generation_nan": int(frame["generation_mwh"].isna().sum()),
        "sampled_values_verified": frame.attrs["sampled_values_verified"],
    }, indent=2))


if __name__ == "__main__":
    main()